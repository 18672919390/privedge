"""Build the team-internal material package: one folder plus one zip with code, data, results and Word docs.

Run with the project venv:
    D:\\比赛\\privedge\\.venv\\Scripts\\python.exe D:\\比赛\\competition_entry\\tools\\build_share_package.py
Outputs go to D:\\比赛\\share\\. Raw downloads, virtual environments, node_modules and caches are never included.
Needs portable pandoc at D:\\比赛\\tmp\\pandoc\\pandoc-3.12\\pandoc.exe and python-docx in D:\\比赛\\tmp\\pydeps_docx
(pip install --target there).
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
TOOLS = Path(__file__).resolve().parent
PANDOC = ROOT / "tmp" / "pandoc" / "pandoc-3.12" / "pandoc.exe"
sys.path.insert(0, str(ROOT / "tmp" / "pydeps_docx"))
import docx  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

DOCS = ROOT / "competition_entry" / "docs" / "v4"
START = TOOLS / "从这里开始.md"
TASKS = DOCS / "09_任务清单（待认领）.md"

INCLUDE = [
    "privedge",
    "competition_entry/docs/v4",
    "competition_entry/tools",
    "review/top-k-share-20260925",
    "dataset_research/TRD_AUDIT_RESULT_20260929.md",
    "dataset_research/runs/olist_brazil",
    "competition_entry/v4_audit/outputs/massive_ldp_recompute_20260929.json",
]
SKIP_DIRS = {".venv", "node_modules", "__pycache__", ".pytest_cache", ".git"}
SKIP_FILES = {"tsconfig.tsbuildinfo"}
SECRET_PATTERN = re.compile(r"(\.env$|\.pem$|\.key$|kaggle\.json$|token|secret|credential|id_rsa)", re.I)

# Math is written as \( \) and \[ \]; "<" in inequalities and "~" in ranges must stay literal text.
PANDOC_FROM = ("markdown+tex_math_single_backslash+lists_without_preceding_blankline+east_asian_line_breaks"
               "-tex_math_dollars-raw_html-smart-subscript-superscript-yaml_metadata_block")
FONT = "微软雅黑"
LANDSCAPE_MIN_COLUMNS = 8

CSV_COLUMNS = ["编号", "类别", "任务", "必做/可选", "做什么", "交付", "怎样算完成", "需要的能力", "人日", "建议截止",
               "依赖", "认领人", "协助人", "备注"]


def collect() -> list[tuple[Path, str]]:
    files = []
    for entry in INCLUDE:
        source = ROOT / entry
        paths = [source] if source.is_file() else sorted(p for p in source.rglob("*") if p.is_file())
        for path in paths:
            rel = path.relative_to(ROOT).as_posix()
            parts = rel.split("/")
            if any(p in SKIP_DIRS for p in parts) or parts[-1] in SKIP_FILES:
                continue
            if SECRET_PATTERN.search(path.name):
                raise SystemExit(f"refusing to package a possible credential file: {rel}")
            files.append((path, rel))
    return files


def _width(text: str) -> int:
    text = re.sub(r"\*\*|`|\\\(|\\\)", "", text)
    return sum(2 if ord(c) > 0x2E80 else 1 for c in text)


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _column_rule(rows: list[list[str]], rule: list[str]) -> str:
    """Rewrite a pipe-table separator so dash counts follow content width; pandoc turns them into column widths."""
    parts = []
    for j, spec in enumerate(rule):
        widths = [_width(r[j]) for r in rows if j < len(r)] or [3]
        mean, longest = sum(widths) / len(widths), max(widths)
        token = max((_width(t) for r in rows if j < len(r) for t in re.split(r"[\s，、；（）]+", r[j])), default=0)
        floor = max(min(widths[0], 8) + 3, min(token, 8) + 3, 6, longest + 3 if longest <= 12 else 0)
        weight = max(floor, round(0.5 * mean + 0.5 * min(longest, 2 * mean)))
        dashes = "-" * min(weight, 45)
        parts.append((":" if spec.startswith(":") else "") + dashes + (":" if spec.endswith(":") else ""))
    return "|" + "|".join(parts) + "|"


def _prepare(text: str) -> str:
    """Make the GitHub-style docs parse cleanly as pandoc Markdown.

    Pandoc needs a blank line before tables, block quotes and headings that follow a paragraph line,
    and nested items indented past the parent marker (the docs use 2 or 3 spaces, one level deep).
    """
    item = re.compile(r"^ *(?:[-*]|\d+\.) ")
    rule = re.compile(r"^\|(\s*:?-+:?\s*\|)+\s*$")
    lines = text.splitlines()
    out: list[str] = []
    fenced = False
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.lstrip().startswith("```"):
            fenced = not fenced
        elif not fenced:
            line = re.sub(r"^ {1,3}(?=(?:[-*]|\d+\.) )", "    ", line)
            prev = out[-1] if out else ""
            starts_block = line.startswith(("|", ">", "#"))
            prev_is_text = prev.strip() and not item.match(prev) and not prev.startswith(("|", ">", "#"))
            if (starts_block and prev_is_text) or (line.startswith("#") and prev.strip()):
                out.append("")
            if line.startswith("|") and i + 1 < len(lines) and rule.match(lines[i + 1]):
                j = i + 2
                while j < len(lines) and lines[j].startswith("|"):
                    j += 1
                rows = [_cells(line)] + [_cells(r) for r in lines[i + 2:j]]
                out += [line, _column_rule(rows, _cells(lines[i + 1]))] + lines[i + 2:j]
                i = j
                continue
        out.append(line)
        i += 1
    return "\n".join(out) + "\n"


def _element(tag: str, children: tuple = (), **attrs: str):
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn(f"w:{key}"), value)
    for child in children:
        el.append(child)
    return el


def _shading(fill: str):
    return _element("w:shd", val="clear", color="auto", fill=fill)


def _chinese_fonts(rfonts) -> None:
    for attr in ("w:eastAsiaTheme", "w:asciiTheme", "w:hAnsiTheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rfonts.set(qn(attr), FONT)


def reference_docx(path: Path) -> None:
    """Pandoc's default reference doc, restyled: Chinese font, A4, bordered tables, compact code."""
    path.write_bytes(subprocess.run([str(PANDOC), "--print-default-data-file", "reference.docx"],
                                    check=True, capture_output=True).stdout)
    document = docx.Document(str(path))
    styles = document.styles
    defaults = styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    _chinese_fonts(defaults.find(qn("w:rFonts")))
    lang = defaults.find(qn("w:lang"))
    if lang is not None:
        lang.set(qn("w:eastAsia"), "zh-CN")
    for rfonts in styles.element.iter(qn("w:rFonts")):
        if rfonts.get(qn("w:asciiTheme")) or rfonts.get(qn("w:eastAsiaTheme")):
            _chinese_fonts(rfonts)
    styles["Normal"].font.size = Pt(10.5)
    for name, size in {"Title": 20, "Heading 1": 16, "Heading 2": 14, "Heading 3": 12, "Heading 4": 11}.items():
        font = styles[name].font
        font.size, font.bold, font.color.rgb = Pt(size), True, RGBColor(0x1F, 0x38, 0x64)
    code_style = styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    code_style.base_style = styles["Normal"]
    code_style.paragraph_format.space_before = code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1.0
    for name in ("Source Code", "Verbatim Char"):
        style = styles[name]
        style.font.name, style.font.size = "Consolas", Pt(9)
    styles["Source Code"].element.get_or_add_pPr().insert_element_before(
        _shading("F5F5F5"), "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
        "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
        "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
        "w:rPr", "w:sectPr", "w:pPrChange")
    caption = styles["Image Caption"]
    caption.font.size, caption.font.italic = Pt(9), False
    caption.font.color.rgb = RGBColor(0x59, 0x63, 0x6E)
    caption.paragraph_format.alignment = 1
    styles["Captioned Figure"].paragraph_format.alignment = 1
    for name in ("Body Text", "First Paragraph", "Compact"):
        styles[name].paragraph_format.line_spacing = 1.2
    tbl_pr = styles["Table"].element.find(qn("w:tblPr"))
    for tag in ("w:tblBorders", "w:tblCellMar"):
        for old in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(old)
    edges = [_element(f"w:{e}", val="single", sz="4", space="0", color="BFBFBF")
             for e in ("top", "left", "bottom", "right", "insideH", "insideV")]
    tbl_pr.insert_element_before(_element("w:tblBorders", edges), "w:shd", "w:tblLayout", "w:tblCellMar",
                                 "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
    margins = [_element(f"w:{e}", w=str(v), type="dxa")
               for e, v in (("top", 30), ("left", 80), ("bottom", 30), ("right", 80))]
    tbl_pr.insert_element_before(_element("w:tblCellMar", margins), "w:tblLook", "w:tblCaption",
                                 "w:tblDescription", "w:tblPrChange")
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2)
    document.save(str(path))


def _polish(path: Path, landscape: bool) -> None:
    document = docx.Document(str(path))
    for table in document.tables:
        for r, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            for tag in ("w:cantSplit",) + (("w:tblHeader",) if r == 0 else ()):
                if tr_pr.find(qn(tag)) is None:
                    tr_pr.append(OxmlElement(tag))
            for cell in row.cells:
                if r == 0:
                    cell._tc.get_or_add_tcPr().insert_element_before(
                        _shading("EEF2F7"), "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
                        "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.1
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        if r == 0:
                            run.font.bold = True
    if landscape:
        for section in document.sections:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = Cm(29.7), Cm(21)
            section.left_margin = section.right_margin = Cm(1.6)
            section.top_margin = section.bottom_margin = Cm(1.8)
    document.save(str(path))


def to_docx(md_path: Path, reference: Path, work: Path) -> bytes:
    text = md_path.read_text(encoding="utf-8")
    landscape = max([line.count("|") - 1 for line in text.splitlines() if line.startswith("|")] or [0]) >= LANDSCAPE_MIN_COLUMNS
    source, target = work / "in.md", work / "out.docx"
    source.write_text(_prepare(text), encoding="utf-8")
    subprocess.run([str(PANDOC), str(source), "-f", PANDOC_FROM, "-t", "docx", "-o", str(target),
                    "--reference-doc", str(reference), "--resource-path", str(md_path.parent), "--columns", "10"],
                   check=True)
    _polish(target, landscape)
    return target.read_bytes()


def word_docs(work: Path) -> dict[str, bytes]:
    reference = work / "reference.docx"
    reference_docx(reference)
    out = {f"文档/{d.stem}.docx": to_docx(d, reference, work) for d in sorted(DOCS.glob("*.md"))}
    out["文档/privedge_README.docx"] = to_docx(ROOT / "privedge" / "README.md", reference, work)
    out["从这里开始.docx"] = to_docx(START, reference, work)
    return out


def task_csv() -> bytes:
    """Flatten every task row of the 09 task list into one sheet (UTF-8 with BOM so Excel opens it)."""
    rows, header, category = [], None, ""
    for line in TASKS.read_text(encoding="utf-8").splitlines():
        if line.startswith("### "):
            category, header = line[4:].strip(), None
        elif line.startswith("|") and not re.match(r"^\|[\s\-|]+\|$", line):
            cells = [re.sub(r"\*\*|`", "", c) for c in _cells(line)]
            if cells[0] == "编号":
                header = cells
            elif header and re.fullmatch(r"[A-DX]\d+", cells[0]):
                record = dict(zip(header, cells))
                record["类别"] = category
                record.setdefault("建议截止", record.get("截止", ""))
                rows.append(record)
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=CSV_COLUMNS, extrasaction="ignore", lineterminator="\r\n")
    writer.writeheader()
    writer.writerows(rows)
    if len(rows) < 30:
        raise SystemExit(f"only {len(rows)} tasks parsed from {TASKS.name}")
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--name", default="PrivEdge_资料")
    args = parser.parse_args()
    out_dir = ROOT / "share"
    out_dir.mkdir(exist_ok=True)
    folder, archive_path = out_dir / args.name, out_dir / f"{args.name}.zip"
    if folder.exists():
        shutil.rmtree(folder)

    entries: dict[str, bytes] = {rel: path.read_bytes() for path, rel in collect()}
    with tempfile.TemporaryDirectory() as work:
        entries.update(word_docs(Path(work)))
    entries["任务认领表.csv"] = task_csv()
    entries["文件清单.txt"] = ("\n".join(f"{sha256(data)}  {rel}" for rel, data in sorted(entries.items()))
                           + "\n").encode("utf-8")

    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for rel, data in entries.items():
            target = folder / rel
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(data)
            archive.writestr(f"{args.name}/{rel}", data)
    with zipfile.ZipFile(archive_path) as archive:
        if archive.testzip():
            raise SystemExit("CRC check failed")
    archive_path.with_suffix(".zip.sha256").write_text(f"{sha256(archive_path.read_bytes())}  {archive_path.name}\n",
                                                       encoding="utf-8")
    print(f"{folder}\n{archive_path}: {archive_path.stat().st_size / 1e6:.1f} MB, {len(entries)} files")


if __name__ == "__main__":
    main()
