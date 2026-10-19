"""Create the E34 operator manual from frozen, observed program behavior."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'competition_entry/docs/v4/PrivEdge_完整演示与复现操作手册_E34_20261018.docx'
IMAGES=ROOT/'competition_entry/docs/v4/backend_validation/20261017_review_delivery/browser'
doc=Document()
s=doc.sections[0]
s.page_width,s.page_height=Cm(21),Cm(29.7)
s.top_margin,s.bottom_margin=Cm(1.8),Cm(1.7)
s.left_margin,s.right_margin=Cm(2),Cm(2)
s.footer_distance=Cm(.8)
for name,size in [('Normal',11),('Title',22),('Heading 1',17),('Heading 2',12),('Caption',9)]:
    st=doc.styles[name]
    st.font.name='微软雅黑';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
    st.paragraph_format.line_spacing=Pt(16.5 if name=='Normal' else size*1.3)
    st.paragraph_format.space_after=Pt(5)
    pp=st._element.get_or_add_pPr()
    grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');pp.append(grid)
    for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
    if name.startswith('Heading'):
        st.paragraph_format.space_before=Pt(8)
        st.paragraph_format.keep_with_next=True
doc.styles['Title'].paragraph_format.space_after=Pt(12)
foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('PrivEdge E34  操作手册  ·  ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)

def p(text,bold=False,size=None):
    x=doc.add_paragraph();r=x.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    return x
def h(text):doc.add_heading(text,2)
def page(title):
    if len(doc.paragraphs)>0:doc.add_page_break()
    doc.add_heading(title,1)
def step(n,text):p(f'{n}. {text}')
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Cm(width)
    for i,(c,txt) in enumerate(zip(t.rows[0].cells,headers)):c.text=txt;c.width=Cm(widths[i])
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for vals in rows:
        for i,(c,txt) in enumerate(zip(t.add_row().cells,vals)):c.text=str(txt);c.width=Cm(widths[i])
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            tc.append(b)
            m=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'95');e.set(qn('w:type'),'dxa');m.append(e)
            tc.append(m)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF5');tc.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=Pt(14)
                for run in para.runs:run.font.size=Pt(10);run.bold=ri==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
def picture(name,caption,width=16.5):
    para=doc.add_paragraph();para.paragraph_format.keep_with_next=True
    para.paragraph_format.line_spacing=1.0
    para.add_run().add_picture(str(IMAGES/name),width=Cm(width))
    doc.add_paragraph(caption,style='Caption')
def say(text):p('现场讲解：'+text)
def code(text):
    x=p(text,size=9);x.paragraph_format.line_spacing=1.1
    for r in x.runs:r.font.name='Consolas'

doc.add_paragraph('PrivEdge 完整演示与复现操作手册',style='Title')
p('适用版本 E34  ·  Windows x64  ·  2026年10月18日',size=10)
p('照着本手册操作，可以启动程序，完整演示正常交换、两类攻击和账本拒答，展示历史实验，再从随包标准数据复现实验流程。每一步都说明点哪里、应看到什么、怎样讲解，以及应保存哪些记录。')
p('演示主线：多源 LDP 摘要的抗投毒可信交换。新区域热门品类 Top-10 推断是应用例子。',True)
h('先选你要完成的路线')
table(['路线','按哪些页操作','完成后得到什么'],[
('现场完整演示','第2至8页，约10至15分钟的排练路线','五页功能、真实处理事件、攻击对比和账本拒答'),
('功能补充与记录','第9至10页','十类攻击巡检、另一数据集、重放与JSON证据'),
('离线实验复现','第11至12页，耗时取决于电脑','新实验目录、复核结果、形式化结果和重绘图表'),
('排错与上场前检查','第13至14页','问题定位、简短讲稿和逐项验收表')],[3.2,6.1,7.7])
h('本机直接使用的位置')
p('演示程序文件夹：',True)
code(r'D:\比赛\share\PrivEdge_Windows_x64_E34_20261017T213814')
p('进入这个文件夹后双击“启动程序.cmd”。本机已有完整文件夹，不必再解压。换电脑时使用完整程序ZIP，整体解压；不要只复制启动脚本。')
p('开发源码在 D:\\比赛\\privedge；当前报告和本手册在 D:\\比赛\\competition_entry\\docs\\v4。现场演示优先使用上面的冻结程序文件夹，避免开发改动造成版本混用。')
h('什么叫复现完成')
p('现场单次结果、历史统计存档、快速流程核验、完整旧参数重算是四种不同工作。本手册逐一标明。能打开实验中心不等于重新运行实验；快速216行核验也不能替代报告中的完整统计。当前严格确认性A8仍未完成，既有四方法结果按描述性证据使用。')

page('1 启动并核对程序版本')
h('正式演示前先做一次自检')
step(1,'打开第1页的演示程序文件夹，确认同时存在 runtime、application、启动程序.cmd、验证程序.cmd、运行测试.cmd、先读我_使用指南.html。')
step(2,'双击“验证程序.cmd”，等待结束。它核对6064个冻结载荷文件，并执行三数据预测、证据读取和账本副本检查；输出中 integrity.ok、functional.ok 应为 true。窗口可滚动查看完整输出。')
step(3,'双击“运行测试.cmd”，等待测试汇总。此冻结版本实测为239 passed，另有3条既有极小ε数值警告。实际失败、异常退出或文件缺失须先排查，不能直接沿用旧通过记录。')
step(4,'双击“启动程序.cmd”，保持启动窗口打开。浏览器通常自动打开；若没有，把窗口中“PrivEdge:”后面的完整网址复制到浏览器。默认尝试8765端口，占用时自动使用空闲端口，以窗口实际网址为准。')
step(5,'页面应显示总览、冷启动预测、攻防实验室、账本审计、实验中心五个入口。底部展开“程序 E34 / 0.4.1 · 查看可核对标识”，确认版本和指纹。')
h('当前冻结版本标识')
table(['字段','应核对内容'],[
('程序ZIP','PrivEdge_评审演示与复现_E34_20261017.zip'),
('ZIP大小','66,784,721字节，约66.8 MB'),
('页面版本','E34 / 0.4.1'),
('ZIP文件SHA-256','4c5429d52c847319851d0bd664be76a5a5be1ed11a88243ed3612310a29a44fa'),
('页面源码指纹','4ddec5336d5451bcf27db951eade0c4d9bbc24c289f7aad4b51f36b042b7fc06')],[3.8,13.2])
p('两个指纹范围不同：ZIP指纹代表整个压缩文件；页面源码指纹只覆盖接口说明列出的Python源码、构建网页和配置。两者不应互相比较，也不是第三方安全认证。',size=10)
h('现场注意')
p('程序无需联网安装Python或Node，也无需上传个人数据。所有角色在本机同一进程模拟，只监听127.0.0.1。不要关闭启动窗口；演示结束才按Ctrl+C。浏览器刷新清空页面记录，服务重启才清空内存账本和密钥。')

page('2 先讲清程序完整处理过程')
p('打开“总览”，先看架构，再开始操作。你扮演目标侧网关的操作者：选择数据集和留出的目标区域，让其他来源发布隐私摘要，系统核验后融合出榜单。')
picture('overview.png','E34真实页面的架构图。箭头表示摘要与审计信息的传递，不代表实际跨机器部署。',14)
table(['阶段','程序实际做什么'],[
('用户与发布者','读取随包标准表；抽样并进行OLH本地扰动；发布端筛查伪造报告，估计频率与方差，产生Top-m摘要。'),
('签名与登记','发布者签名；账本在登记前验签并记录窗口预算与摘要内容的哈希值，形成哈希链。'),
('目标网关','检查来源、角色、窗口、格式、签名、登记内容和账本完整性，记录接受或拒收原因。'),
('融合与输出','对通过准入的摘要进行留一截尾、多源支持和权重封顶，给出榜单、三态和审计依据。'),
('离线评价','目标区域真值只用于计算Recall和NDCG，不进入发布、准入或融合。')],[3.1,13.9])
say('“隐私扰动不保证报告诚实，签名也不保证数值真实。我们把发布端筛查、网关协议检查和鲁棒融合串成一个可操作、可核查的交换流程。”')

page('3 演示正常摘要交换')
step(1,'右上角选择“美团外卖（北京11商圈）”。在总览点击“配置正常交换”，程序跳到“冷启动预测”并填写预设；这个按钮只填参数，不会自动运行。')
p('固定参数：目标商圈3；隐私预算ε=1；每区域发布Top-m=20；每源采样用户数上限5000；随机种子20260929。',True)
step(2,'点击“开始预测”。观察本地扰动、签名发布、账本登记、网关接受、融合返回的事件。日志按节奏显示真实执行记录，不表示网络真实延时。')
step(3,'等到出现“上一次已完成结果”及运行编号，再向下看榜单、检查结论、降级原因、离线评价和审计记录。')
picture('prediction.png','固定参数的实测结果：10个类别，DEGRADED，Recall@10为70%，NDCG@10约0.897。窗口、运行编号和哈希每次可能不同。',16.3)
say('“这里没有人为注入攻击，但来源间存在真实差异，因此系统提示降级并保留结果。本次找回了真实前十中的七个；这说明检查状态与预测准确度需要分开看。”')
p('验收点：10个来源参与；有10项输出；固定示例为LOW_AGREEMENT，源间一致度约0.180；Recall为0.7，NDCG原值0.8968。不要为了显示NORMAL反复换种子。耗时不要求与截图一致。',size=10)

page('4 演示合法签名的数值投毒')
step(1,'正常交换完成后，回总览点击“配置权重劫持”；或进入“攻防实验室”，点击“带入上次正常运行参数和种子”。核对仍是TRD、商圈3、ε1、m20、5000、seed20260929。')
step(2,'选择“权重劫持（谎报极小方差 + 注入）”，请求攻击来源数选1，点击“发起攻击”。不要把MGA的β当成此攻击的共同强度。')
step(3,'先看“三层处理结果”：发布端筛查、网关准入、整次融合结论分别展示。本例可能接受全部10个来源，因为攻击者能合法签名；准入未拒收不代表数值没有攻击。')
step(4,'看“本次效果对比（离线评价）”，展开“本次攻击比较：查看数据表”，并比较左右榜单中标出的攻击目标。')
picture('attack_metrics.png','同一请求下的两种处理结果。无防御基线是无防御方差加权，不是简单平均。',16.3)
say('“攻击者用合法身份谎报方差并注入目标类别。无防御方差加权下三个目标都入榜；PrivEdge下仍有一个入榜，影响减小，但并未彻底消失。”')
p('固定实测：AS基线3/3=100%，PrivEdge为1/3=33.3%；Recall为0.400对0.600，NDCG约0.616对0.864。整次为DEGRADED。AS是本次指定目标中的入榜比例，不是多次试验的攻击成功概率。',size=10)

page('5 演示未重新签名的摘要篡改')
step(1,'返回总览点击“配置签名篡改”，或在攻防实验室选择“中间人篡改摘要（不重新签名）”。核对与第4页相同的固定基础参数，攻击来源数为1。')
step(2,'点击“发起攻击”，等结果返回。展开完整攻击与拒收事件，寻找SIGNATURE_INVALID。该场景修改了签名后的内容，没有重新签名，网关应识别不匹配。')
step(3,'核对网关统计：固定实测为接受9个来源、拒收1个。再单独看整次请求结论：实测仍为DEGRADED，而不是REFUSED。')
step(4,'查看AS说明：本场景没有指定注入目标，AS显示“不适用”；这不是0%，也不是所有攻击都被完全防住的统计证明。')
h('把三个层次讲清楚')
table(['观察内容','发生位置','能够说明什么'],[
('过滤报告数','发布端','对用户扰动报告的异常筛查；剔除数不等于已确认伪造数。'),
('SIGNATURE_INVALID','网关准入','这份被篡改摘要未通过验签，不能继续作为合格输入。'),
('DEGRADED或REFUSED','整次决策','由全部合格输入、候选数、一致性与链状态决定，不能仅凭一源拒收推定。')],[4.1,3.2,9.7])
h('与上一场景的差异')
p('权重劫持利用合法签名发布虚假数值，主要观察鲁棒融合；本场景破坏签名内容一致性，主要观察准入拒收。二者不是同一个防御机制，因此界面分别给出来源层事件和整次请求结论。')
say('“签名解决的是谁发布、内容有没有被改；它不能判断一个合法发布者是不是在说真话。上一项靠数值防线限制影响，这一项由网关验签拒收。”')
h('本页应留存的记录')
p('下载本次JSON；截图保留攻击名称、固定参数、SIGNATURE_INVALID、接受9与拒收1，以及整次DEGRADED。不要只截一个红色拒收字样而省掉整次结论。')
p('截图中的来源名称、窗口号、审计哈希可能与本次不同。核对原因码和行为，不要求时间戳或哈希相同。')

page('6 演示账本审计与损坏副本拒答')
step(1,'保持TRD数据集，进入“账本审计”。确认已有业务条目，记录当前条目数、链头、完整性和最新窗口预算。若为空，点“前往冷启动预测”并先运行一次。')
step(2,'找到篡改演练，记录序号输入0，点击“执行演练”。0表示第一条；有效范围为0至条目数减1。程序只修改副本中的ε，不改原始账本。')
step(3,'等待副本演练及原链复核完成：副本应提示链异常，决策为REFUSED，原因LEDGER_CORRUPTED，不输出推荐。')
step(4,'对照演练前后的原链：条目数、链头和完整性应相同。页面完成后会重新读取原链，不能只展示副本拒答而省略原链复核。')
picture('ledger.png','真实演练截图中的60条只是当时累积数量。你的条目数取决于已运行次数，验收要求是本次演练前后原链不变。',14.5)
say('“链上登记内容一旦不一致，系统拒绝基于损坏账本继续推荐。演示使用副本，不会破坏后续操作。它不等于能抵抗宿主把整条链一起重写。”')
p('预算按窗口解释；不同运行建立新窗口，不要把历史窗口所有ε直接相加当成当前窗口预算。跨窗口隐私解释依赖部署假设，账本不能识别用户集合是否重叠。',size=10)

page('7 展示历史实验及证据范围')
step(1,'打开“实验中心”。先指出页面“历史证据库 · 只读”：这里读取已登记结果，点击刷新不触发实验重算。')
step(2,'在“四方法对照 · 描述性结果”卡片选择美团TRD、ε=2、无攻击。读出四种方法：简单平均、无防御方差加权、逐项截尾、PrivEdge。')
step(3,'展开“方法定义与解释范围”和“完整来源与文件指纹”。切换一个攻击场景，先确认条件一致，再比较指标；无攻击AS不适用。需要留存时点击“下载当前条件历史数据 JSON”。')
picture('experiments.png','四方法卡片示例。它与现场攻防页的两方法单次比较不同，不能互相代替。',13.9)
step(4,'向下看“严格预注册绑定核验（A8）”：NOT_RUN与上方描述性结果可用能够并存。随后浏览隐私效用、投毒、权重劫持、形式化、通信开销及预注册文档。')
say('“我们同时保留有利和不利结果。无攻击时不宣称比简单平均更准；有些场景已有鲁棒基线更好。四方法现有结果用于描述性比较，不能当作独立确认性复验。”')
p('关于形式化：结果只覆盖符号协议模型中的性质。模型不自动证明整个Python程序，也不证明LDP的数值性质或统计鲁棒性。',size=10)

page('8 补充展示巡检和跨数据集运行')
h('十类攻击巡检')
step(1,'回TRD攻防实验室，保持商圈3、ε1、m20、5000、seed20260929、攻击源数1。点击“一键全攻击巡检”。MGA使用β=10%，其它攻击不把β当统一强度。')
step(2,'等待进度达到“巡检完成 · 10 / 10”。逐行查看攻击类型、实际来源、准入统计、三态、AS及运行编号，再下载本次记录。不同攻击不应预设同一结论。')
step(3,'若展示取消行为，可另起一次巡检后切页，再返回。应明确显示“未完成”，不会自动续跑；已开始的后端计算可能继续，但客户端停止后续请求。保存完整巡检后再做此项，避免覆盖待导出的页面快照。')
table(['类别','重点检查'],[
('mga 伪造用户','发布端筛查与AS；β为伪造数相对诚实采样数，按实现round取整。'),
('compromised / hijack / demote / negative','合法发布者的摘要数值攻击；看融合影响和残余目标。'),
('tamper / replay / equivocation / impostor / budget_burn','协议异常；看具体拒收事件和整次结论，不能仅靠攻击名称判断成功与否。')],[7.2,9.8])
h('跨数据集的可重复示例')
step(1,'选择“Massive-STEPS城市签到”，进入冷启动预测，设Petaling Jaya@2018、ε=2、Top-m=40、5000用户上限、seed=1，无攻击。')
step(2,'点击开始预测并等结束。E34实测为6个来源、10个命名类别、NORMAL、Recall=0.5、NDCG约0.806；可能同时有SOURCE_ANOMALY附加告警。')
step(3,'对比TRD：同一程序能换数据和类别域；NORMAL只表示当前检查通过，Recall仍可能为50%。这恰好说明三态不能代替准确度指标。')
p('Olist也可选择并运行，目标必须选该数据集下的区域。未在本手册固定的参数不承诺具体数值；记录实际参数和结果即可。切回TRD会保留该数据集已完成的页面快照，刷新网页则清空。')

page('9 重放结果并导出现场证据')
h('用相同参数重放')
step(1,'在冷启动预测或单次攻击已完成结果旁点击“下载本次记录 JSON”，先保存原结果。下载位置由浏览器设置决定，通常在“下载”目录。')
step(2,'点击“使用这次参数和种子再运行”。核对榜单、状态、原因和离线指标与先前一致；窗口、运行ID、请求ID、时间、耗时、签名和审计哈希不要求相同。')
step(3,'如果修改控制区种子但尚未运行，旧结果仍是“上次已完成结果”，不能把它讲成新参数的输出。种子留空的单次请求会随机生成；巡检留空则明确使用20260929。')
h('导出的文件里核对什么')
table(['字段组','用途'],[
('program_identity','E34、0.4.1、源码指纹范围、三套标准数据的文件哈希。'),
('payload.request_snapshot','本次请求的真实参数；不要用后来改动的页面控制值代替。'),
('run_id / window / seed','区分不同执行、预算窗口和随机种子。'),
('decision / baseline / evaluation','本次输出、无防御对照及离线指标。'),
('sources / events / audit','来源明细、处理事件、准入原因、链头与审计记录。')],[6,11])
p('JSON是请求和响应的可核对记录，不是外部认证。现场结果可以下载；历史条件数据和账本快照也有各自的JSON下载入口，文件类型不要混用。')
h('建议建立一个演示记录文件夹')
p('在桌面或另一个可写位置建“PrivEdge_本次演示_日期”。按顺序存正常交换、权重劫持、签名篡改、账本副本、完整巡检及历史条件JSON；截图应带参数与结论。不要把文件放回冻结载荷并改写清单。')
p('本机最终包已验证的正常记录为run-834835633b，源码指纹4ddec533…；这是先前验收记录，现场不要求生成同一编号。新演示必须保存你自己这一次的文件。')
h('结束前保存')
p('切页能够保留已完成快照，但刷新网页会清空页面内存。服务重启会重建密钥、窗口状态和账本。演示结束前先下载，再在启动窗口按Ctrl+C停止；仅关闭浏览器不会停止服务。')

page('10 离线复现快速实验和协议模型')
p('以下操作在程序文件夹根目录进行，不在网页里运行。建议现场讲解前完成一次，把结果索引准备好；演示时可展示真实记录。每次执行会先核对冻结载荷，再写新的会话与运行目录。')
h('快速核心统计复现')
step(1,'双击“快速核心统计复现.cmd”，保持窗口打开，等待三套数据计算及独立复核完成。')
step(2,'这次使用各数据集首个目标、旧种子101、ε=1与2、m=20、每源5000用户、9个场景、4种方法：每数据72行，合计216行。')
step(3,'窗口应以PASSED结束并给出“结果索引.html”路径。打开对应的新会话，逐项检查三数据和review状态。保留smoke_test=true标记；这是工程烟测，不是报告的完整统计复验。')
h('重跑形式化验证')
step(1,'双击“重跑形式化验证.cmd”。程序使用随包ProVerif 2.05，不需另装工具。')
step(2,'在新会话中打开结果索引和原始输出。E34既有模型中Q1至Q5性质成立；S1与S2的查询不成立，以及四个负对照中的预定性质失败，是预期结果，不能只数TRUE来判成功。')
step(3,'先看会话整体PASSED，再对照具体查询和负对照解释。模型结果支持协议逻辑，不能扩展成整个实现没有漏洞。')
h('从历史证据重建图表')
step(1,'双击“从历史证据重建图表.cmd”。它读取历史结果，生成5组实验图与2张系统图，共7组PNG/SVG、14个文件。')
step(2,'在本次会话figures目录查看图表。它只重新绘图，没有重新计算原统计结果；不同电脑字体、图片字节或绘图时间可能不同。')
h('完成判定')
p('E34包在本机已实际完成上述三条流程。你在另一台电脑运行时，仍以本次PASSED、输出文件及日志为准。FAILED或INTERRUPTED须保留日志，不能把上一次成功目录当作本次成功。')
p('这些入口不会把新结果自动登记到实验中心，也不会覆盖旧实验或原始数据。程序能从随包标准表重算；Olist/Massive从外部原始数据重建的预处理链仍不完整。')

page('11 完整重算的范围与结果位置')
h('完整A8四方法旧参数重算')
step(1,'需要重算完整四方法网格时，双击“完整A8四方法重算.cmd”。它遍历三数据集全部目标、5个既有种子101/211/307/401/503、ε1/2、m20、每源5000、9场景与4方法，最后独立复核。')
step(2,'这是较长计算，不建议在短时间讲解中临时启动。E34本轮验收核对了完整命令计划，但没有重新执行全部网格；不能将已提供按钮表述为本轮全量重算已完成。')
step(3,'旧种子重算用于复现既有描述性结果，不能补成新的确认性A8。若要进行新实验或改变算法、参数、判定规则，应另写预注册；结果差也保留。')
h('会话和原始结果的相对位置')
table(['相对程序根目录的路径','内容'],[
('reproduction_outputs\\模式-时间-标识\\结果索引.html','本次执行的总入口：状态、每步日志和结果链接。'),
('同一会话的session.json与command_plan.json','实际状态、参数、命令、输出位置和核验记录。'),
('同一会话的logs\\','stdout及stderr；失败先看这里。'),
('application\\runs\\实验名\\新运行ID\\','新算法或形式化原始结果，包括manifest和相应CSV/JSON。'),
('同一会话的review\\或figures\\','独立汇总复核、比较图或历史重绘产物。')],[9.2,7.8])
h('只查看计划而不运行')
p('在程序文件夹空白处打开PowerShell，复制下面任一行。--plan只显示命令，不计算统计、不创建复现会话。')
code(r'.\runtime\python.exe -B -X utf8 .\run_reproduction.py --mode quick --plan')
code(r'.\runtime\python.exe -B -X utf8 .\run_reproduction.py --mode full --plan')
code(r'.\runtime\python.exe -B -X utf8 .\run_reproduction.py --mode formal --plan')
code(r'.\runtime\python.exe -B -X utf8 .\run_reproduction.py --mode figures --plan')
h('报告其他实验怎样复现')
p('完整A8入口不包含所有历史消融、效用、开销和报告排版。根目录“先读我_使用指南.html”第5节及“复现索引.json”按E编号列出了security_v1、security_v3、utility_v1、negative_check和overhead_v1的模块、参数与结果位置。按对应条目执行，不把A8完成等同于整份报告所有实验已重跑。')

page('12 常见现象与处理办法')
table(['遇到的现象','先做什么与怎样解释'],[
('网页打不开或连接失败','检查启动窗口是否仍开着；复制它这次显示的实际网址。端口可能变化，旧书签不一定有效。'),
('双击后一闪而过或提示缺文件','确认整体解压、路径可读写、runtime和application未被单独移动；保存错误。必要时从原ZIP重新解压到新目录。'),
('校验提示某文件哈希不同','记录错误文件列表，使用未修改的冻结包重试；不要删检查或改manifest把错误“修成通过”。'),
('账本为空或演练按钮禁用','先在当前数据集完成一次预测；检查记录序号是否为0至条目数减1的整数。'),
('无攻击却显示降级','先读具体原因。固定TRD示例为来源一致性低；它仍可返回榜单，不是程序崩溃。'),
('权重劫持没有签名拒收','这是合法发布者签名的数值投毒；检查融合前后AS和排序，不能用“是否拒收”判断全部效果。'),
('签名篡改只拒一源而整次不拒答','来源准入与整次决策是不同层级；其余合格来源仍能产生结果。'),
('AS显示不适用或未提供','该场景没有定义注入目标或字段缺失；不要补成0，不要解释成攻击完全失败。'),
('改了参数但榜单没变','检查是否尚未执行新请求；上次快照不会随控制区改参而变。需要点击运行并等新编号。'),
('切页后巡检只剩部分行','本轮已停止后续请求。读取“未完成”状态，重新发起新巡检，不声称10项已完成。'),
('相同种子但审计哈希不同','预期现象。窗口、密钥、签名、ID和时间可变；比较数值结果和参数，不能要求所有字节一致。'),
('复现失败或图表缺中文','保留本次stderr与结果索引；核对实际字体环境。重跑产生新目录，不覆盖失败记录。')],[5.1,11.9])
p('界面刷新会清掉页面快照，但不会清服务账本；重启服务会清账本。账本损坏路径应空榜拒答；候选不足引起的REFUSED可能仍展示部分候选，要按原因解释。',size=10)

page('13 上场讲稿与完成检查')
h('按这个顺序讲约十分钟')
table(['环节','可以直接说的内容'],[
('开场与总览','“我们做的是多源LDP摘要的抗投毒可信交换。新区域选品是展示用例。隐私机制与签名各有边界，因此需要分层核验和融合。”'),
('正常交换','“现在现场生成扰动报告、签名摘要、预算登记和榜单。来源分歧使本次降级；准确度另看Recall和NDCG。”'),
('合法数值投毒','“合法签名不能保证数值真实。本次攻击目标由基线3/3降到PrivEdge的1/3，但仍有残余风险。”'),
('未重签篡改','“这一类内容改动由签名核验识别，拒收一份来源。其余来源能继续产生结果，单份拒收不等于整次拒答。”'),
('账本演练','“只篡改副本就会触发链损坏拒答；原链条目、链头和完整性不变。”'),
('历史与复现','“实验中心读取存档，四方法结论是描述性的。程序另有真实重算入口和结果索引，便于独立核查。”')],[3.4,13.6])
h('演示结束前逐项确认')
p('□ 使用E34冻结程序，页面版本与指纹核对正确。\n□ 正常、权重劫持、签名篡改三个场景都实际运行并保存JSON。\n□ 展示了三层处理、AS分子分母、准确度与三态的区别。\n□ 账本副本拒答，原链复核不变；没有破坏原始数据。\n□ 历史证据与现场单次结果分开讲，严格A8状态如实保留。\n□ 若演示巡检，确认10/10后再称完整；中断不算完成。\n□ 换电脑时按本手册再启动、自检、操作、导出并停止；保留实际记录。')
h('结束语')
say('“当前作品把隐私摘要生成、分层防御、签名与预算核验、结果解释和证据复现连成了完整原型。保证有明确前提；它不是多机器生产部署，也不保证无攻击更准或所有攻击都失败。”')
p('核对依据：E34冻结程序指南与复现索引、2026年10月17日真实浏览器验收、最终包启动与下载记录、包内239项回归及三条实际离线复现记录。本手册新增说明，不修改冻结程序、历史实验或原始数据。',size=9)

doc.core_properties.title='PrivEdge 完整演示与复现操作手册'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.core_properties.subject='E34 操作步骤 现场讲解 结果核对 离线复现'
for st in doc.styles:
    pp=st._element.find(qn('w:pPr'))
    if pp is not None:
        for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
for para in list(doc.paragraphs)+[p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]:
    pp=para._p.get_or_add_pPr()
    grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');pp.append(grid)
    for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
for shape in doc.inline_shapes:
    shape._inline.docPr.set('descr','PrivEdge E34实际验收页面截图，具体参数与解释见邻近图注')
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(json.dumps({'output':str(OUT),'planned_pages':14,'pictures':len(doc.inline_shapes),'tables':len(doc.tables)},ensure_ascii=False))
