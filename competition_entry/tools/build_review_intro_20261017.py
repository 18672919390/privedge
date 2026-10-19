"""Build the anonymous one-page introduction from reviewed project facts."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'competition_entry/docs/v4/PrivEdge_作品简介_20261017.docx'

PARAGRAPHS = [
    ('解决的问题', '多方希望共享有用的统计摘要，又不愿暴露用户个人数据。本地差分隐私可以保护个人输入，却不能阻止伪造报告；数字签名可以确认发布来源，也不能保证合法发布者提交的数值诚实。PrivEdge 面向这一缺口，构建多源 LDP 摘要的抗投毒可信交换原型。'),
    ('方案设计', '系统按各环节能够访问的信息设置两层防御。用户在本地扰动类别值，发布端从扰动报告中筛查异常联合支持，生成带签名的 Top-m 摘要；目标网关仅利用多源摘要，通过支持规则、留一截尾和权重封顶限制恶意来源的影响。签名、主体角色、发布窗口及预算登记共同绑定摘要的发布与使用，账本损坏时拒绝继续提供推荐。'),
    ('主要成果', '作品实现了从隐私摘要生成、发布登记、网关准入到融合与审计的完整流程，并给出支持规则和数值影响的条件性分析。在 TRD 的既有五种子实验中，ε=1、每源发布20项、单源伪造数为诚实报告数10%时，筛查与截尾将攻击目标入榜比例由90.30%降至0%。三套数据的四方法描述性对照同时保留部分鲁棒基线更优和双源合谋仍有残余风险的结果；无攻击时不宣称比简单平均更准确。实验条件、统计口径与来源见作品报告第三章及证据索引。'),
    ('程序体验', '新区域热门品类推断是应用例子。Windows x64 程序包内置运行环境和三套标准数据，提供总览、冷启动预测、攻防实验室、账本审计和实验中心。使用者可以执行摘要交换、十类攻击及账本副本篡改演练，查看榜单、准入原因和审计记录，并核对归档实验结果。'),
    ('验证范围', '当前系统在同一进程模拟各角色；检查通过表示来源与一致性检查达标，不代表预测准确或所有攻击均被阻止。理论影响界依赖留一诚实多数等前提，四方法结果按描述性证据解释。程序以随包标准表为复现起点，实际部署与从公开原始数据重建仍需另行验证。'),
]

def set_font(run, size=12, bold=False):
    run.font.name = '宋体'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')

def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.1)
    sec.left_margin = sec.right_margin = Cm(2.35)
    style = doc.styles['Normal']
    style.font.name, style.font.size = '宋体', Pt(12)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    set_font(p.add_run('PrivEdge 作品简介'), 18, True)
    p = doc.add_paragraph()
    set_font(p.add_run('作品名称：'), bold=True)
    set_font(p.add_run('PrivEdge 多源 LDP 摘要的抗投毒可信交换'))
    p.paragraph_format.space_after = Pt(10)
    for label, body in PARAGRAPHS:
        p = doc.add_paragraph()
        p.paragraph_format.widow_control = True
        set_font(p.add_run(label + '：'), bold=True)
        set_font(p.add_run(body))
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = 'PrivEdge 作品简介'
    doc.core_properties.subject = '多源 LDP 摘要的抗投毒可信交换'
    doc.core_properties.comments = ''
    for tree in (doc._element, doc.styles._element):
        for border in tree.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    main()
