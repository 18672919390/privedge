"""Build the review-organized report without modifying the E33 source.

Usage: bundled Python rewrite_report_review_20261017.py [--config JSON] [--out-dir DIR]
The optional config supplies current validated interface prose and screenshot paths;
without it the report uses verified E33 capabilities. No experiment is run here.
"""
from __future__ import annotations
import argparse, copy, hashlib, json, re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / 'competition_entry/docs/v4/PrivEdge_作品报告_前后端最终整合版_20261017.docx'
TMP = ROOT / 'tmp/report_review_revision_20261017'
EXPECTED_BASE = 'b961d96bf81f45d3125e189eae202e7bee68531144d8b498128a64f61e0a19eb'
REPORT_NAME = 'PrivEdge_作品报告_评审整编版_20261017.docx'
HISTORY_NAME = 'PrivEdge_工程验证与历史记录附件_20261017.docx'
ABSTRACT = '多方共享本地差分隐私（LDP）统计摘要时，扰动机制保护个人输入，却不能阻止伪造报告；数字签名确认发布来源，也不能保证合法发布者提交的数值诚实。PrivEdge面向这一缺口，实现多源LDP摘要的抗投毒可信交换。系统按信息可见性设置两层防御：发布端从扰动报告中筛查异常联合支持，网关仅基于Top-m摘要组合多源支持、留一截尾与权重封顶；签名、窗口、角色和预算登记绑定摘要的发布与实际使用，链损坏触发拒答。以新区域品类推断为例，作品在三套数据上验证统计效用与攻击影响，并以协议模型和程序测试核对交换行为。TRD中ε=1、m=20、单源伪造数为诚实数10%的场景下，既有五种子实验的攻击目标入榜比例由0.9030降至0；四方法描述性对照同时保留无攻击无普遍优势、部分鲁棒基线更优及双源合谋残余风险。系统输出检查通过、降级或拒答及原因，检查通过不代表预测准确。理论影响界依赖留一诚实多数等条件，原型验证不替代真实分布式部署。'

def text_of(el):
    return ''.join(el.xpath('.//w:t/text()'))

def put(parent, name, attrs=None):
    el=OxmlElement(name)
    for k,v in (attrs or {}).items(): el.set(qn(k),str(v))
    parent.append(el); return el

class Builder:
    def __init__(self, config, history=False):
        self.doc=Document(BASE); self.original=[copy.deepcopy(x) for x in self.doc._element.body]
        self.body=self.doc._element.body
        for x in list(self.body): self.body.remove(x)
        self.cfg=config; self.history=history; self.headings=[]; self.copied=[]; self.figure_keys=[]
        self.body.append(copy.deepcopy(self.original[-1]))
    def add(self,el):
        self.body.insert(len(self.body)-1,el); return el
    def p(self,text='',bold=False,center=False,keep=False):
        p=self.doc.add_paragraph(text)
        pf=p.paragraph_format;pf.line_spacing=1.5;pf.space_after=Pt(6);pf.space_before=Pt(0)
        pf.first_line_indent=Pt(24) if text else Pt(0);pf.keep_with_next=keep
        if center: p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Pt(0)
        for r in p.runs:
            r.font.name='宋体';r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体');r.font.size=Pt(12);r.bold=bold
        return p
    def h(self,text,level=2,page=False):
        style={1:'Heading 1',2:'PE Heading 2',3:'PE Heading 3'}.get(level,'PE Heading 3')
        # The template's custom styles have localized display names; IDs are stable.
        sid={1:'2',2:'PEHeading2',3:'PEHeading3'}[level]
        p=self.p(text);p._p.get_or_add_pPr().get_or_add_pStyle().set(qn('w:val'),sid)
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
        p.paragraph_format.space_before=Pt(12 if level>1 else 0);p.paragraph_format.space_after=Pt(8)
        p.paragraph_format.page_break_before=page
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:r.bold=True;r.font.size=Pt({1:16,2:14,3:12}[level])
        key=f'RR_{len(self.headings)+1:03}'
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(1000+len(self.headings)));start.set(qn('w:name'),key)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(1000+len(self.headings)))
        p._p.insert(1,start);p._p.append(end);self.headings.append({'anchor':key,'text':text,'level':level})
        return p
    def copy(self,a,b=None,replace=None):
        for i in range(a,(a if b is None else b)+1):
            e=copy.deepcopy(self.original[i]);self.copied.append(i)
            for n in list(e.xpath('.//w:bookmarkStart|.//w:bookmarkEnd')):n.getparent().remove(n)
            for n in list(e.xpath('.//w:sectPr')):n.getparent().remove(n)
            if e.tag==qn('w:sectPr'):continue
            if e.tag==qn('w:p') and not text_of(e) and not e.xpath('.//w:drawing|.//m:oMath'):continue
            if replace:
                for n in e.xpath('.//w:t'):
                    for old,new in replace.items():n.text=(n.text or '').replace(old,new)
            self.add(e)
    def heading_copy(self,i,new=None,level=3,page=False):
        return self.h(new or text_of(self.original[i]),level,page)
    def table(self,headers,rows,widths=None):
        t=self.doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        widths=widths or [6.15/len(headers)]*len(headers)
        for i,s in enumerate(headers):t.rows[0].cells[i].text=s
        for row in rows:
            cells=t.add_row().cells
            for i,s in enumerate(row):cells[i].text=str(s)
        for row in t.rows:
            for i,c in enumerate(row.cells):c.width=Inches(widths[i])
        for i,c in enumerate(t.columns):c.width=Inches(widths[i])
        return t
    def cap(self,text):
        p=self.p(text,center=True,keep=True);p.paragraph_format.first_line_indent=Pt(0);return p
    def fig(self,key,index,caption,note=''):
        caption=self.cfg.get('figure_captions',{}).get(key,caption)
        path=self.cfg.get('figures',{}).get(key)
        if path:
            path=Path(path)
            if not path.is_absolute():path=ROOT/path
            p=self.p(center=True,keep=True);p.add_run().add_picture(str(path),width=Inches(6.1))
            self.figure_keys.append(key)
        else:self.copy(index)
        p=self.cap(caption);p.paragraph_format.keep_with_next=bool(note)
        if note:self.p(note)
    def section_paragraph(self,index):self.add(copy.deepcopy(self.original[index]))
    def front(self):
        for i in range(23):self.add(copy.deepcopy(self.original[i]))
        self.p('目     录',bold=True,center=True);self.toc_marker=self.p('')
        self.section_paragraph(88)
        self.h('摘要',1);self.p(ABSTRACT);self.p('关键词：本地差分隐私；多源摘要；数据投毒防御；可信交换；可审计原型')
        self.section_paragraph(92)
    def toc(self, skip=0):
        marker=self.toc_marker._p
        for item in self.headings[skip:]:
            p=OxmlElement('w:p');pp=put(p,'w:pPr');put(pp,'w:spacing',{'w:line':360,'w:lineRule':'auto','w:after':20})
            put(pp,'w:ind',{'w:left':(item['level']-1)*240});tabs=put(pp,'w:tabs');put(tabs,'w:tab',{'w:val':'right','w:leader':'dot','w:pos':8820})
            hyp=put(p,'w:hyperlink',{'w:anchor':item['anchor'],'w:history':1})
            for value in (item['text'],'\t',str(self.cfg.get('toc_pages',{}).get(item['anchor'],''))):
                r=put(hyp,'w:r');rp=put(r,'w:rPr');put(rp,'w:rFonts',{'w:ascii':'宋体','w:hAnsi':'宋体','w:eastAsia':'宋体'});put(rp,'w:sz',{'w:val':24});put(rp,'w:color',{'w:val':'000000'})
                if value=='\t':put(r,'w:tab')
                else:put(r,'w:t').text=value
            marker.addprevious(p)
        marker.getparent().remove(marker)
    def normalize(self):
        for p in self.doc.paragraphs:
            if p._p.xpath('.//w:sectPr'):continue
            pp=p._p.get_or_add_pPr()
            # Preserve mathematics and figures; body text inherits official 12pt Songti.
            if not p._p.xpath('.//w:drawing'):
                sp=pp.find(qn('w:spacing'))
                if sp is None:sp=put(pp,'w:spacing')
                sp.set(qn('w:line'),'360');sp.set(qn('w:lineRule'),'auto')
            style=pp.find(qn('w:pStyle'));sid=style.get(qn('w:val')) if style is not None else ''
            if sid in ('1','Normal'):
                if style is not None:pp.remove(style)
            if sid not in ('2','PEHeading2','PEHeading3') and not p._p.xpath('.//w:drawing|.//m:oMath'):
                # Preserve the cover's title sizing, otherwise body and captions are 12pt.
                if list(self.body).index(p._p)>22 or self.history:
                    for r in p.runs:
                        r.font.size=Pt(12);r.font.name='宋体';r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        for t in self.doc.tables:
            pr=t._tbl.tblPr
            for n in list(pr.findall(qn('w:tblpPr'))):pr.remove(n)
            for n in list(pr.findall(qn('w:tblBorders'))):pr.remove(n)
            borders=put(pr,'w:tblBorders')
            for side in ('top','left','bottom','right','insideH','insideV'):put(borders,'w:'+side,{'w:val':'single','w:sz':4,'w:color':'D9D9D9'})
            for j,row in enumerate(t.rows):
                trp=row._tr.get_or_add_trPr()
                for n in list(trp.findall(qn('w:trHeight'))):trp.remove(n)
                if trp.find(qn('w:cantSplit')) is None:put(trp,'w:cantSplit')
                if j==0 and trp.find(qn('w:tblHeader')) is None:put(trp,'w:tblHeader')
                for cell in row.cells:
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tcpr=cell._tc.get_or_add_tcPr()
                    if j==0:
                        shd=tcpr.find(qn('w:shd'))
                        if shd is None:shd=put(tcpr,'w:shd')
                        shd.set(qn('w:fill'),'E7E6E6')
                    margins=tcpr.find(qn('w:tcMar'))
                    if margins is None:margins=put(tcpr,'w:tcMar')
                    for side in ('top','bottom','left','right'):
                        e=margins.find(qn('w:'+side))
                        if e is None:e=put(margins,'w:'+side)
                        e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa')
                    for p in cell.paragraphs:
                        p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=(j==0)
                        for r in p.runs:r.font.size=Pt(12);r.font.name='宋体';r.bold=j==0;r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        cp=self.doc.core_properties;cp.author='';cp.last_modified_by='';cp.comments='';cp.title='PrivEdge工程验证与历史记录附件' if self.history else 'PrivEdge多源LDP摘要的抗投毒可信交换';cp.subject='';cp.keywords=''
        for part in self.doc.part.package.parts:
            if str(part.partname)=='/docProps/app.xml':
                blob=re.sub(rb'<Company>.*?</Company>',b'<Company/>',part.blob)
                blob=re.sub(rb'<Application>.*?</Application>',b'<Application/>',blob)
                # These retained-template counters are stale until a Word host
                # recalculates them; final PDF pagination is checked separately.
                part._blob=re.sub(rb'<(Pages|Words|Characters|CharactersWithSpaces|Lines|Paragraphs|TotalTime)>.*?</\1>',b'',blob)
        upd=self.doc.settings.element.find(qn('w:updateFields'))
        if upd is not None:self.doc.settings.element.remove(upd)
        # Plain-text math inherited from the draft contains literal underscores
        # and combining accents. Typeset the explicit mathematical whitelist as
        # native inline math; do not touch API identifiers or the 15 original
        # display objects. The whitelist preserves each symbol's meaning.
        tokens={
            'ᾱ':('accent','α','̄'), 'C̃_i':('accent_sub','C','̃','i'),
            'S_i^H':('subsup','S','i','H'),
            **{f'{x}_i':('sub',x,'i') for x in ('c','M','R','V','U','A','S')},
            **{f'{x}_{{{sub}}}':('sub',x,sub) for x,sub in
               [('m','−s,i'),('d','−s,i'),('η','−s,i'),('z','s,i'),
                ('v','s,i'),('z','h,i'),('v','h,i'),('z','b,i'),('Σ','h∈H')]},
            **{f'{x}_{sub}':('sub',x,sub) for x,sub in
               [('u','s'),('ω','s'),('ω','h'),('ω','b'),('u','h'),('α','F'),('A','t'),('H','a')]},
            '10^-12':('sup','10','−12'),'e^ε':('sup','e','ε'),
        }
        pattern='('+'|'.join(re.escape(k) for k in sorted(tokens,key=len,reverse=True))+')'
        def math_run(holder,value):
            mr=put(holder,'m:r');put(mr,'m:t').text=value
        def math_symbol(token):
            spec=tokens[token];kind=spec[0];math=OxmlElement('m:oMath')
            if kind in ('accent','accent_sub'):
                holder=math
                if kind=='accent_sub':
                    sub=put(math,'m:sSub');holder=put(sub,'m:e')
                acc=put(holder,'m:acc');pr=put(acc,'m:accPr');put(pr,'m:chr',{'m:val':spec[2]})
                math_run(put(acc,'m:e'),spec[1])
                if kind=='accent_sub':math_run(put(sub,'m:sub'),spec[3])
            else:
                obj=put(math,{'sub':'m:sSub','sup':'m:sSup','subsup':'m:sSubSup'}[kind])
                math_run(put(obj,'m:e'),spec[1])
                if kind in ('sub','subsup'):math_run(put(obj,'m:sub'),spec[2])
                if kind in ('sup','subsup'):math_run(put(obj,'m:sup'),spec[3] if kind=='subsup' else spec[2])
            return math
        for r in list(self.doc._element.xpath('.//w:r[w:t]')):
            raw=''.join(r.xpath('./w:t/text()'))
            if not re.search(pattern,raw):continue
            parent=r.getparent();pos=parent.index(r)
            for part in re.split(pattern,raw):
                if not part:continue
                if part in tokens:
                    parent.insert(pos,math_symbol(part))
                else:
                    nr=copy.deepcopy(r)
                    for child in list(nr):
                        if child.tag!=qn('w:rPr'):nr.remove(child)
                    t=put(nr,'w:t');t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=part;parent.insert(pos,nr)
                pos+=1
            parent.remove(r)
    def save(self,path):
        self.normalize();self.doc.save(path)
        return {'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'bytes':path.stat().st_size,'headings':self.headings,'copied_source_blocks':self.copied,'omml':len(self.doc._element.xpath('.//m:oMath')),'tables':len(self.doc.tables),'inline_shapes':len(self.doc.inline_shapes)}

def build_report(cfg):
    b=Builder(cfg);b.front()
    b.h('第一章 作品概述',1)
    b.h('1.1 研究背景')
    b.p('多个组织希望共享统计规律，却不愿直接交换个人行为记录。本地差分隐私允许用户先扰动类别再上报，但也使接收方难以区分诚实随机化和刻意构造的报告。与此同时，持有合法密钥的发布者可以生成验签正确的虚假摘要。隐私保护、来源认证和统计可信性因此需要分别处理。')
    b.p('PrivEdge解决的具体问题是：目标网关只接收多个来源的稀疏Top-m摘要时，如何限制恶意来源对候选品类和融合分数的影响，并说明哪些摘要被使用、哪些检查失败。作品以多源LDP摘要的抗投毒可信交换为主线，新区域热门品类推断只是检验这条交换流程的应用例子。')
    b.p('设计的关键约束是信息不对称。发布端能查看用户级扰动报告，可以筛查异常共同支持；网关无法回看这些报告，只能利用摘要间的支持关系与数值分布。系统据此在两处设置防线，并用签名登记和实际使用内容绑定连接两层。最终输出检查结论、Top-k及审计原因，离线预测指标另行评价。')
    b.h('1.2 相关工作')
    b.p('相关研究分别解决隐私估计、扰动报告投毒、鲁棒聚合或协议认证。下面按它们能够观察的信息和提供的保证定位本作品，既注明继承的方法，也说明组合设计的新增范围。')
    b.heading_copy(99);b.copy(100);b.copy(105)
    b.heading_copy(106);b.copy(107)
    b.p('本项目使用MGA风格的有界随机搜索构造伪造报告，具体攻击预算和比例见3.1.2。发布端筛查和网关摘要级防御处理不同输入：前者不能约束被攻破发布者直接伪造摘要，后者无法回溯用户级报告。')
    b.heading_copy(111,'1.2.3 arXiv 2506.02156与恶意报告检测')
    b.copy(112);b.p('本作品没有实现该论文的Diffstats、ASD或RSN，也没有把其效用恢复结论当作本队结果。当前发布端采用联合支持项集启发式，具体过程及阈值限制见2.5.2。该论文是相关研究参照，作品的实验数据、程序实现和影响界以本报告对应证据为准。')
    b.heading_copy(122);b.copy(123,124)
    b.p('PrivEdge另组合原始来源支持、留一截尾与权重封顶，其适用条件与分数界见2.5.3、2.6.2及附录A。')
    b.heading_copy(128);b.copy(129,130)
    b.cap('表1-1 相关方法与PrivEdge的关系')
    b.table(['研究方向','已有方法及输入','本作品使用与新增范围'],[
        ['LDP估计','OLH等[1–3]；用户级随机报告','沿用OLH结构；将稀疏摘要接入带来源、窗口和预算登记的交换流程。'],
        ['报告投毒检测','MGA与频繁项集思路[4]；Diffstats、ASD、RSN[5]','参考威胁及统计检测思想；实现本地联合支持启发式，不声称复现上述检测器。'],
        ['鲁棒聚合','坐标截尾与Krum[6,7]；完整梯度向量','对照采用缺失补零的截尾适配；本作品组合原始支持、留一截尾和权重封顶。'],
        ['审计与协议','隐私组合[8]、ProVerif[9]、透明日志[10]','约束签名登记与实际使用；内存哈希链不等同分布式透明日志。']], [1.05,2.0,3.1])
    b.h('1.3 作品贡献')
    b.p('第一，按可见信息分层设置抗投毒防线。发布端针对扰动报告的异常共同支持进行筛查；网关针对合法签名下的摘要作恶限制跨源影响，两者覆盖不同攻击入口。')
    b.p('第二，同时约束候选支持与数值影响。当恶意来源数f<t时，原始Top-m支持门槛排除从未被诚实来源发布的新增项；留一截尾和权重封顶约束极端频率与谎报方差。给出各自的条件性保证，并用四方法对照保留代价和失效场景。')
    b.p('第三，将签名登记与使用内容绑定，并输出可解释失败。注册角色、签名窗口、预算与账本摘要共同参与准入；链损坏进入拒答，单份摘要拒收保留原因。符号模型、程序检查和页面证据分别支撑协议逻辑与实现行为。')
    b.p('三项贡献的具体机制在第二章展开，统计和系统证据在第三章报告，第四章归纳与已有能力的差异及适用范围。')
    b.h('1.4 应用场景')
    b.p('以新区域品类推断为例，已有区域分别发布经过LDP扰动的品类摘要，新区域网关在不接收个人原始类别的前提下形成Top-10建议。离线评价使用目标区域真实分布衡量排序质量，目标真值不输入预测。此例可说明摘要交换的可行性；选品收益、库存和地域迁移效果尚需真实业务试验。')

    b.h('第二章 作品设计与实现',1,True)
    b.heading_copy(138,level=2);b.copy(139,141)
    b.heading_copy(142,level=2)
    b.p('系统区分伪造用户、被攻破发布者和网络消息攻击者，并明确注册表、账本及宿主的信任假设。签名认证和统计防御分别应对不同能力，不能互相替代。')
    b.copy(144,145)
    b.h('2.3 设计目标')
    b.p('设计目标不是在所有数据上提高无攻击精度，而是在保留可用统计效用的同时，对不同攻击入口建立可检查的防线。下表给出从威胁到证据的对应关系，后文沿同一路径说明实现。')
    b.cap('表2-1 威胁 可见信息 防线与验证位置')
    b.table(['威胁或条件','处理方可见信息','处理与输出','验证位置'],[
        ['伪造用户报告','发布端可见OLH报告','联合支持筛查后重新估计','2.5.2；3.4.2'],
        ['合法密钥伪造频率或方差','网关仅见多源Top-m摘要','支持门槛、留一截尾、封顶权重','2.5.3；2.6.2；3.4.1'],
        ['篡改、冒名、旧窗口','签名、角色注册与登记摘要','拒收不合格摘要并保留原因','2.5.4；3.4.4'],
        ['链损坏或可用来源不足','账本链、准入集与候选集','拒答；不生成伪推荐','2.5.5；3.4.5'],
        ['隐私与实际效用','固定参与集合与离线目标真值','ε-LDP论证；独立评价排序','2.6.1；3.4.3']], [1.3,1.55,2.0,1.3])
    b.heading_copy(148,level=2);b.copy(149,153)
    b.heading_copy(154,level=2);b.copy(155)
    for hi,end in [(156,171),(172,182),(183,199),(200,208)]:
        b.heading_copy(hi);b.copy(hi+1,end)
    b.h('2.5.5 三态输出与解释',3)
    for s in cfg.get('state_prose',[]):b.p(s)
    if not cfg.get('state_prose'):b.copy(210,213)
    b.heading_copy(214,level=2);b.copy(215)
    for hi,end in [(216,221),(222,236)]:
        b.heading_copy(hi);b.copy(hi+1,end)
    b.heading_copy(237);b.copy(238,240)
    b.copy(241,replace={'当前表述沿用已登记模型结果，没有在本轮重复运行ProVerif。':cfg.get('formal_rerun_note','当前表述沿用已登记模型结果，没有在本轮重复运行ProVerif。')})
    b.h('2.6.4 设计约束与实现验证',3)
    b.p('安全论证与实现检查采用不同证据。隐私性依赖固定参与集合和随机化机制；支持及影响界依赖来源身份与留一诚实多数；ProVerif只讨论抽象协议事件。程序测试再检查模型未覆盖的数值输入、消息解析、重复来源、账本链和资源路径，避免用任一层的通过代替另一层保证。')
    b.p('当前实现对非法ε、非有限数值及结构错误返回可解释错误，拒绝无效登记；预测前先验证账本，链损坏时直接拒答。WebSocket字段错误不会启动计算，并允许连接继续接收后续合法请求。静态资源限制在构建目录内。上述约束均有保留的回归用例，当前版本执行结果见3.4.5。')
    b.p('历史缺陷、修复前失败现场及历次候选验证另存《工程验证与历史记录附件》。这些记录用于追溯检查如何形成；本报告不把历史电脑的通过记录继承为当前程序的异机验收，也不将软件测试数量视为算法安全强度。')
    b.h('2.7 前端设计与实现')
    b.p('五个页面围绕同一条使用路径组织：先理解交换架构，再运行预测，观察攻击影响，核对账本，最后查看统计证据。前端只显示后端返回的决策与已登记统计，不在浏览器重算算法均值、区间或隐私预算。')
    b.copy(252)
    b.h('2.7.1 总览与冷启动预测',3)
    for s in cfg.get('overview_prose',[]):b.p(s)
    if not cfg.get('overview_prose'):b.copy(254)
    b.fig('overview',255,'图2-5 总览与多源摘要可信交换架构',cfg.get('figure_notes',{}).get('overview','来源图是角色关系说明；当前原型在同一宿主进程中模拟这些角色。'))
    for s in cfg.get('prediction_prose',[]):b.p(s)
    if not cfg.get('prediction_prose'):b.copy(259)
    b.fig('prediction',260,'图2-6 冷启动预测结果与检查结论',cfg.get('figure_notes',{}).get('prediction',text_of(b.original[262])))
    b.copy(263)
    b.h('2.7.2 攻防实验室与账本审计',3)
    for s in cfg.get('attack_prose',[]):b.p(s)
    if not cfg.get('attack_prose'):b.copy(265,266)
    b.fig('attack',267,'图2-7 攻击条件与结果比较',cfg.get('figure_notes',{}).get('attack',text_of(b.original[269])))
    if cfg.get('figures',{}).get('attack_metrics'):
        b.fig('attack_metrics',267,'图2-7(b) 同一次攻击的指标与分子分母',cfg.get('figure_notes',{}).get('attack_metrics',''))
    b.copy(270)
    b.fig('ledger',271,'图2-8 账本完整性与副本篡改演练',cfg.get('figure_notes',{}).get('ledger',text_of(b.original[273])))
    b.copy(274)
    b.h('2.7.3 实验中心与证据展示',3)
    for s in cfg.get('experiment_prose',[]):b.p(s)
    if not cfg.get('experiment_prose'):b.copy(276)
    b.fig('experiments',277,'图2-9 四方法对照与证据状态',cfg.get('figure_notes',{}).get('experiments','TRD、ε=2、无攻击。四方法描述性比较可用，严格绑定入口无登记；两种状态分别表达，页面没有生成新的统计实验。'))
    b.copy(280)
    b.p('四方法结果属于描述性比较。界面展示其冻结时序和种子复用限制，统计文件核验通过只表示读取产物与登记一致。单次演示、页面核对与工程回归均不会把描述性证据升级为确认性结论。')
    b.heading_copy(282,level=2);b.copy(283)
    b.heading_copy(284);b.copy(286)
    for row in cfg.get('extra_api_rows',[]):
        cells=b.doc.tables[-1].add_row().cells
        for i,value in enumerate(row):cells[i].text=value
    b.copy(287,288)
    for s in cfg.get('backend_additions',[]):b.p(s)
    b.heading_copy(289);b.copy(290,295,replace={'E28补齐原件并核验后，E29以独立描述性入口接入；':'描述性结果通过独立入口接入；'})

    b.h('第三章 作品测试与分析',1,True)
    b.p('本章回答三个问题：防御相对简单平均及鲁棒适配基线有何收益和代价；支持与权重设计在什么条件下失效；协议和程序是否执行了设计规定的检查。统计结果来自既有登记运行及独立复算，当前程序验证单独报告。单次界面运行不并入五种子统计。')
    b.h('3.1 测试方案')
    b.h('3.1.1 系统功能与接口测试',3);b.copy(304);b.cap('表3-1 系统测试层次与预期行为');b.copy(306)
    b.h('3.1.2 算法实验与对照设计',3)
    b.copy(310,replace={'本轮对已有算法实验进行来源核验和独立统计复算，没有重新生成扰动报告，也没有完整重跑算法网格。':'正文统计结果采用已有算法实验，经来源核验与独立统计复算保留。E34另执行小规模离线重跑以检查复现入口，不代替原完整算法网格，也不并入下列统计。'})
    b.copy(311);b.copy(108)
    b.p('四方法比较把简单平均、无防御方差加权、缺失补零的坐标截尾适配和PrivEdge置于相同数据与攻击条件下。前三种方法共享未筛查摘要；PrivEdge包含发布端筛查与摘要级防御，是整套统计方案比较。对照的完整配置与解释边界见3.4.1。')
    b.copy(312,replace={'3.4.7':'3.4.1','四方法原始运行与完整复核已补齐（E28）':'四方法原始运行与复核记录登记为E28'})
    b.h('3.1.3 程序验收与复现范围',3)
    b.p('程序验收检查解压完整性、服务启动、五页交互、错误恢复及自动回归。当前程序从新解压目录启动，运行时、依赖和源码路径均须位于程序包内。该检查证明本机解压运行的可行性；最终冻结程序仍需第二台实体电脑验收。历史B候选的macOS与Windows记录保存在工程附件，不能自动转移为新版本结论。')
    b.h('3.2 测试环境')
    if cfg.get('environment_prose'):b.p(cfg['environment_prose'])
    else:b.copy(319)
    b.cap('表3-2 统计与当前程序验证环境')
    b.table(['验证对象','环境与范围','结论边界'],[
        ['登记算法运行','Windows 11；Python 3.12.14；NumPy 2.3.5。数据、种子、代码指纹由运行manifest记录。','对应既有数据与攻击网格；本次复算不重新生成报告。'],
        ['当前整合程序',cfg.get('environment','本机Windows x64；隔离CPython及锁定依赖；现成前端由本地服务托管。'),'同机新解压与自动检查；不代表所有设备兼容或真实分布式部署。']], [1.2,3.1,1.85])
    b.h('3.3 数据集与评价标准');b.copy(326)
    b.heading_copy(327);b.cap('表3-3 三套标准数据规模');b.copy(329,332)
    b.heading_copy(333);b.copy(334,338)
    b.h('3.4 结果展示与分析')
    b.h('3.4.1 四方法描述性对照',3)
    b.copy(414,replace={'三组原始运行已补交并登记E28':'三组原始运行登记为E28'});b.copy(416)
    b.p('以下先比较无攻击代价，再比较同一ε下的双源权重劫持。完整两档ε、九场景共216个方法单元保存在描述性证据入口，未展示切片不视为缺失。表内95%经验区间先在目标内平均五个种子，再对目标作2,000次bootstrap；目标共享来源，未作多重比较校正。')
    b.cap('表3-4 无攻击对照  ε=2  均值与目标级95%经验区间');b.copy(419);b.copy(420)
    b.cap('表3-5 双源权重劫持对照  ε=2  均值与目标级95%经验区间');b.copy(422,423)
    b.p('该切片也显示鲁棒适配基线更优的情况。TRD中坐标截尾的AS为0.0182，低于PrivEdge的0.0606，NDCG为0.9608与0.9596；Olist中AS为0.1667与0.1889，PrivEdge的NDCG仅略高（0.9385对0.9378）。因此不能把“防御组件更多”解释成统计指标全面领先。')
    b.p('这些结果支持的结论是：无防御方差加权易受低方差伪报影响，截尾和封顶可以降低部分影响，但支持门槛无法消除达到门槛的合谋。PrivEdge同时实现摘要准入、预算登记和原因输出，属于另一组工程能力；协议能力不能作为统计指标落后的抵销项。该对照直接调用统计决策入口，未执行真实签名或账本登记，协议效果另见3.4.4。')
    b.h('3.4.2 分层防线与权重封顶',3)
    b.p('本节进一步区分防御组件。原方差加权（D0）与发布端筛查＋网关截尾及支持门槛（D3）用于观察两层防线；在D3之上增加权重封顶的配置记为D4。下文的配置名只用于对应原运行记录，不代表不同产品版本。')
    b.copy(383,385,replace={'表3-8':'表3-6'});b.copy(387,390);b.copy(391)
    b.copy(395,398,replace={'图3-3':'图3-1','图3-4':'图3-2'})
    b.h('3.4.3 隐私效用与参数敏感性',3)
    b.p('本节UTILITY_V1与3.4.1的COMPARE_V1是不同运行器。二者都把ε在各自参数网格中的索引写入随机种子序列，但网格分别为[0.5,1,2,4]与[1,2]；因此相同基础种子及相同ε不表示同一扰动实例。表3-7的TRD、ε=2、m=20的Recall为0.8345，表3-4中PrivEdge为0.8127，各自对应原运行，不能跨运行器逐值合并或当作同次结果冲突。')
    b.copy(365,372,replace={'表3-6':'表3-7','图3-1':'图3-3'})
    b.copy(374,380,replace={'表3-7':'表3-8','图3-2':'图3-4'})
    b.h('3.4.4 协议异常路径与模型验证',3)
    b.p('签名与账本验证针对消息在交换过程中的身份、内容和使用条件。它不识别已签名数值是否诚实，因此与3.4.1、3.4.2的统计攻击实验分开验收。下表列出关键异常及预期处置。')
    b.cap('表3-9 协议异常与实现处置')
    b.table(['异常输入','预期行为','证据与解释'],[
        ['未注册密钥、错误角色或签名','登记或准入拒绝','注册表与签名回归；不接受冒名登记。'],
        ['旧窗口、类别域不一致','拒收对应摘要','窗口与域检查；仍允许复用有效窗口摘要。'],
        ['登记摘要与使用内容不一致','拒收对应摘要','摘要digest绑定；合法密钥也不能绕过该一致性检查。'],
        ['预算超限或非法ε','拒绝登记或输入','数值回归；不同窗口重复用户仍属部署假设。'],
        ['账本副本首、中、末条损坏','REFUSED并给出链损坏原因','预测前完整性检查；副本演练不破坏原账本。'],
        ['同源多份有效摘要','拒收该源全部摘要','剩余来源是否足够由决策层判断，不必然整次拒答。']], [1.7,1.55,2.9])
    b.p(cfg.get('protocol_validation_prose','ProVerif 2.05登记的认证、窗口和账本绑定性质在模型内成立；可达性查询证明协议并非永不接受。删除窗口签名、角色绑定、账本核对或登记前验签后，对应负对照破坏预期性质。具体事件和假设见2.6.3；预算数值、哈希链宿主重写和统计防御不由该符号模型覆盖。'))
    b.h('3.4.5 当前程序验证',3)
    for s in cfg.get('validation_prose',[]):b.p(s)
    if not cfg.get('validation_prose'):b.copy(343,344)
    b.p('测试组之间存在重复覆盖，分别报告而不相加为独立样本。极小正ε边界保留3条数值警告，对应用例返回可解释拒答。以上自动检查和本机页面验证不代替最终版本的实体异机验收，也不将文件一致解释为原创权利已确认。')
    b.h('3.4.6 完整应用流程',3)
    for s in cfg.get('application_prose',[]):b.p(s)
    if not cfg.get('application_prose'):b.copy(408,409)
    b.p('该流程把配置、真实运行、结果解释和来源证据连在一起。第二章五个页面的界面图绑定各自参数及运行记录，仅说明功能能够按所示路径执行。攻击示例即使输出检查通过，也可能有攻击目标进入榜单；用户需同时阅读状态、离线评价和攻击指标。当前验证支持本地演示，不代表真实经营收益。')
    b.h('3.4.7 旧管线通信与计算开销',3);b.copy(400,406,replace={'表3-9':'表3-10'})

    b.h('第四章 创新性说明',1,True)
    b.h('4.1 按可见信息分层的两层抗投毒')
    b.p('本作品的系统设计起点是两种攻击入口掌握的信息不同。发布端可访问用户级OLH报告，网关只接收每个来源的Top-m摘要。将同一检测器直接搬到网关既没有所需报告，也无法防止持合法密钥的发布者自行编造摘要。PrivEdge据此把报告级筛查与摘要级约束放在不同信任边界。')
    b.p('发布端用共同支持异常筛除疑似伪造报告并重新估计；网关组合多源支持、留一截尾和权重封顶，处理绕过发布端的摘要作恶。两层通过带来源、窗口和预算的摘要接口连接。新增之处在于这一面向稀疏LDP摘要的信息分工和完整实现，不是重新提出OLH、MGA、坐标截尾或相关论文检测器。')
    b.p('证据闭环由3.4.2的报告攻击与来源攻击实验构成：所测TRD单源10%伪造场景中AS从0.9030降至0，权重劫持复验说明只约束频率仍需约束权重。该结论限定于现有攻击族和参数；四方法对照中已有鲁棒适配基线在部分切片更优，不能写成对所有已有方案的优势。')
    b.h('4.2 多源支持与数值影响约束')
    b.p('稀疏Top-m摘要既包含数值，也包含“哪些来源愿意发布此项”的支持信息。PrivEdge在截尾前统计原始支持，避免后续数值处理被误作来源支持；同时对频率和方差形成的权重分别约束。这样把候选能否进入与进入后影响多大分开分析。')
    b.p('定理1给出t−f个诚实支持的组合界；定理2在留一集合仍有诚实严格多数、无上下文路径及正有限权重等条件下，给出当前融合分数相对其诚实分量的上界。证明完整列于附录A，接口检查负责落实不同来源、有限数值和非空摘要等可检查条件。')
    b.p('两类结论均有明确边界。支持门槛不固定已有类别名次；影响界中的诚实分量仍由当前受攻击输入定义，不是无攻击反事实输出。n=10、f=1、κ=2且诚实源权重相等（W=w）时，恶意权重份额上界为2/11≈0.1818。3.4.1保留Massive双源AS约0.8476以及Olist五源双攻击不满足定理2条件的结果，使理论保证与经验观察相互区分。')
    b.h('4.3 签名登记与使用绑定及可解释失败')
    b.p('仅对摘要签名只能证明密钥来源，无法回答摘要是否按预算登记、是否属于允许窗口，以及预测实际使用的内容是否与登记内容相同。PrivEdge把主体角色、窗口、类别域与摘要内容放入签名和登记核对链，在目标服务预测前检查账本完整性，再把准入结果交给统计决策。')
    b.p('该设计将错误变成可追溯的业务结果：不合格摘要被拒收并记录原因，剩余来源不足或链损坏时拒答，能够继续使用的结果同时保留告警。页面将三态、离线效用和审计分区，避免把安全流程通过混同于预测准确。2.6.3的符号正负对照、3.4.4的异常路径和3.4.5的当前程序验证分别支撑逻辑与实现。')
    b.p('贡献限于原型内可核对的交换流程。内存账本尚无独立见证、防日志分叉与持久化恢复；具备合法权限的宿主重写整条链不在其保证内。四方法统计工具绕过真实协议，不以其输出替代本节的协议验证。')

    b.h('第五章 总结',1,True)
    b.h('5.1 已完成能力')
    b.p('PrivEdge已形成从用户本地扰动、发布端筛查、摘要签名登记，到网关准入、抗投毒融合和结果解释的可运行原型。三套标准数据、五个页面、后端服务及登记实验在同一程序中连通，实验中心可以从报告结论追溯到条件、方法、种子和原始产物。')
    b.p('现有证据支持有限场景中的防御效果与工程可行性，而不支持无攻击普遍提精度或阻止全部合谋的主张。四方法对照将无攻击代价、鲁棒基线较优切片和双源残余攻击同时列出；支持与影响定理只在明确条件下成立。作品将这些边界作为结果解释的一部分，检查通过始终只表示流程条件达标。')
    b.h('5.2 技术局限与后续研究')
    b.p('统计方面，发布端联合支持阈值尚未作独立校准，Top-m缺失、来源异质性和自适应攻击仍会影响防御。当前bootstrap区间反映有限目标集合，不能外推为任意地区的总体保证。后续应使用独立时间窗口与新攻击策略，同时比较效用、误报、漏报和拒答代价。')
    b.p('工程方面，角色仍在同一宿主进程模拟，内存账本与密钥重启后重建；生产使用需要持久化、访问控制、密钥治理和独立审计。跨来源、跨窗口是否重复包含同一用户无法由当前报告条数识别，隐私预算归属需要真实参与管理。')
    b.p('复现方面，本包的起点是随包标准表，完整原始下载与预处理依赖链尚未独立重建。最终冻结版本仍需实体第二电脑验收；成员权利、贡献和正式声明须本人确认。以上未完成事项在附录中明确标识，不由既有测试通过或文件哈希一致替代。')
    b.p('新区域选品展示了不直接交换个人原始记录时的协作方式。后续业务试点应在明确的数据使用条件下另行评价地域差异、选品收益与部署成本，当前排序代理实验不承担这些业务结论。')
    b.h('参考文献',1,True);b.copy(460,469)

    b.h('附录A 完整证明与条件核对',1,True)
    b.copy(472,479,replace={'这一条件是旧稿中需要收紧的地方：':'较弱条件'})
    b.cap('表A-1 证明条件与实现对应');b.copy(481)
    b.p('补充：OLH概率与估计符号。下列等式在理想均匀哈希碰撞模型下给出p、r、q及反偏置计数口径；固定哈希种子的ε-LDP证明不要求该均匀碰撞假设。')
    b.copy(102,104)
    b.h('附录B 结论与证据定位',1,True)
    b.p('本索引只保留阅读正文所需的主要结论入口。每个编号对应不可覆盖的原运行与核对产物；全历史条目、修复失败现场和跨机回传清单见《工程验证与历史记录附件》。统计复算核对已有输出，不证明历史运行前的冻结时序。')
    b.cap('表B-1 正文结论与证据入口')
    b.table(['正文结论','登记证据','读取与解释范围'],[
        ['两层防御的攻击变化','E9；3.4.2','SECURITY_V2原CSV及summary；五种子和全部目标；保留未达预注册条件。'],
        ['隐私效用和参数','E10、E14；3.4.3','UTILITY_V1及封顶等价记录；拒答计0，不能外推业务收益。'],
        ['权重封顶复验','E13；3.4.2','SECURITY_V4新种子，同数据及攻击族；不是新外部数据集。'],
        ['四方法相对表现','E28；3.4.1','三组compare_v1原件与review；8640行、216方法单元；描述性比较。'],
        ['符号协议性质','E16；2.6.3','ProVerif模型、查询、可达性与四个负对照；不含预算数值及统计防御。'],
        ['旧管线通信开销','E6；3.4.7','m=20、合成Zipf数据；旧单机微基准。'],
        ['当前程序验证',cfg.get('evidence_id','E33')+'；3.4.5',cfg.get('evidence_path','backend_validation/20261017_frontend_final')+'；测试、请求、截图与程序指纹。']], [1.45,1.15,3.55])
    b.p('四方法证据由GET /api/evidence/a8/descriptive读取，严格预注册绑定入口单独保留。表格、图与页面均读取规范摘要，并核原件与派生物指纹；哈希一致不补足冻结缺口。既有算法主运行的4,176个统计字段、E28的35,815个字段复算均无差异，细目见工程附件。')
    b.h('附录C 当前程序运行与复现',1,True)
    b.h('C.1 程序版本与三步启动',2)
    b.p(cfg.get('package_prose',text_of(b.original[507]))).alignment=WD_ALIGN_PARAGRAPH.LEFT
    b.p('第一步，将Windows x64程序包完整解压到本地可写目录，保留目录结构。第二步，运行“验证程序.cmd”核对载荷指纹。第三步，运行“启动程序.cmd”，进入其打开的本机页面。便携包包含隔离Python和锁定依赖，运行现成页面无需安装Python或Node.js。')
    b.p('运行“运行测试.cmd”可复查包内自动回归。服务只监听本机；关闭后重新启动会重建当前内存账本与密钥。源码开发与便携运行是不同入口：需要修改源代码时按照程序README及依赖锁重建，修改前端后重新构建静态文件。')
    b.h('C.2 统计复现的起点与限制',2)
    b.p('统计复现以第三章列出的三套标准表为起点，先核对domain.json所记录的数据哈希，再按对应预注册方案在新运行目录执行；不得覆盖旧runs、review或原始数据。每次保留参数、种子、代码和数据指纹、环境、时间及产物清单。')
    for s in cfg.get('repro_prose',[]):b.p(s)
    if not cfg.get('repro_prose'):
        b.p('复算须核对总行数、键、重复次数、状态及全部指标。四方法比较仍按描述性证据解释，不能补写旧manifest取得确认性资格。ProVerif复验需2.05版求解器，已有符号输出不等于本机重跑。完整原始下载和预处理依赖链尚未独立重建，不能仅据脚本存在认定原始数据至结果已经复现。')
    b.p('最终第二电脑验收须记录ZIP指纹、实体设备、系统版本、测试日志、三数据五页操作及失败复测，并由实际执行人确认。历史运行命令、安装和异机回传记录见工程附件；同机新目录解压不代替实体异机验收。')
    b.h('附录D 来源与第三方许可',1,True)
    b.h('D.1 原创范围与代码来源',2)
    b.p('本队自有原始研究代码的原样内核与新增工程层分开登记。OLH、MGA和Ed25519等既有方法按实际使用引用，不列为本队独创；文件一致只证明副本一致，不独立证明作者和权利归属。')
    b.cap('表D-1 实现来源与本项目工作');b.copy(549,replace={'本轮B工作':'验证与复现'})
    b.copy(544,replace={'本轮已核对OLH移植上游为':'已核对的OLH移植上游为'})
    b.p('成员实际贡献、权利归属及原创性声明须本人确认和签署。AI辅助参与代码修改、验证脚本、证据核对与报告整理，不能替代成员的实际贡献或独立人工验收。正式声明另按当届模板提交。').paragraph_format.keep_together=True
    b.h('D.2 数据来源与许可',2);b.cap('表D-2 数据来源与许可核实状态');b.copy(556);b.copy(559)
    b.p('Massive-STEPS另涉及Foursquare Open Source Places，当前代码许可和CC0元数据不能确认全部冻结输入。已有Olist、Massive和TRD研究来源文件的定位及指纹记录，但完整下载处理链与标准表重建尚未独立验证。')
    b.h('D.3 分发组件与未完成确认',2)
    for s in cfg.get('license_prose',[]):b.p(s)
    if not cfg.get('license_prose'):
        b.p('软件组件、运行时和前端依赖分别保留第三方许可及NOTICE；当前程序载荷以对应FILE_MANIFEST.json为准，许可证材料见程序THIRD_PARTY_LICENSES及licenses目录。完整历史31包、npm组件、官方许可响应和119项许可材料核对清单转列工程附件，不把历史安装树等同当前分发物。')
    b.p('仍需本人确认成员作者与贡献、TRD上游授权依据及Massive冻结输入链。已定位的研究脚本不再视为本机缺失，但脚本存在尚不足以证明完整原始输入链或权利归属。最终分发应按实际载荷核对许可，不把运行通过作为授权证明。')
    # Normalize cross-references and chronological design figure numbering.
    for p in b.doc.paragraphs:
        for t in p._p.xpath('.//w:t'):
            s=t.text or ''
            s=s.replace('图2-3  从设备端','图2-2  从设备端').replace('图2-4  网关','图2-3  网关').replace('图2-2 签名发布','图2-4 签名发布')
            s=s.replace('表2-1  容错定理','表2-2  容错定理')
            t.text=s
    b.toc();return b

def build_history(cfg):
    history_cfg=dict(cfg);history_cfg['toc_pages']=cfg.get('history_toc_pages',{})
    b=Builder(history_cfg,True)
    # One numbered section with the same official page and body style.
    sec=b.body[-1]
    put(sec,'w:footerReference',{'r:id':'rId4','w:type':'default'});put(sec,'w:pgNumType',{'w:start':1})
    b.h('PrivEdge 工程验证与历史记录附件',1)
    b.p('编制日期：2026年10月17日',center=True)
    b.p('本附件保留已发生的修复、版本冻结、跨机回传、来源核查和复算记录。其用途是追溯工程过程，不作为当前整合程序的全部验收结论。主报告以《PrivEdge作品报告 评审整编版》为准；历史记录内的“本轮”“当前”均限定在各条所列日期和版本。')
    b.p('文内章节、表号和证据号保留E33原编号，只用于历史定位，不对应评审整编版的新目录。涉及的学术文献见主报告参考文献；后续工具与复现能力以当前主报告及其验证记录为准。')
    b.p('目 录',bold=True,center=True);b.toc_marker=b.p('')
    b.h('一 历次实现检查与修复记录',1,True);b.copy(243,249);b.copy(296)
    b.h('二 历史候选环境与验收过程',1,True);b.copy(301,302);b.copy(308);b.copy(314,324)
    b.h('三 历史功能结果与失败现场',1,True);b.copy(346,363);b.copy(392,394);b.copy(445,446)
    b.h('四 完整证据登记与原件索引',1,True);b.copy(415);b.copy(483,505)
    b.h('五 历史运行命令与异机记录',1,True);b.copy(508,542)
    b.h('六 来源与软件许可核查原记录',1,True);b.copy(544,570)
    b.p('以上记录按E33主报告原段落保存。后续新版本的验证记录应单独追加，不覆写历史产物，不把旧版测试身份或截图继承为新版通过。')
    b.toc(skip=1)
    return b

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--config');ap.add_argument('--out-dir',default=str(TMP/'draft'));ap.add_argument('--report-only',action='store_true');a=ap.parse_args()
    assert hashlib.sha256(BASE.read_bytes()).hexdigest()==EXPECTED_BASE,'E33 reference changed'
    cfg=json.loads(Path(a.config).read_text('utf-8-sig')) if a.config else {}
    out=Path(a.out_dir);out.mkdir(parents=True,exist_ok=True);TMP.mkdir(parents=True,exist_ok=True)
    main=build_report(cfg);result={'report':main.save(out/REPORT_NAME),'abstract_chars':len(ABSTRACT),'config':cfg}
    assert result['report']['omml']>=15,'All original native equations must remain'
    assert len(ABSTRACT)<=500
    if not a.report_only:result['history']=build_history(cfg).save(out/HISTORY_NAME)
    (out/'build_manifest.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),'utf8')
    print(json.dumps({k:{x:v for x,v in val.items() if x in ('path','sha256','bytes','omml','tables','inline_shapes')} if isinstance(val,dict) else val for k,val in result.items() if k!='config'},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
